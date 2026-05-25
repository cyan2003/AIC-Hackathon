"""
Ingestion API endpoints.

POST /ingest/resume — Upload and ingest a resume (PDF/DOCX)
POST /ingest/jd    — Upload and ingest a job description (PDF/DOCX)
"""

from __future__ import annotations

from fastapi import APIRouter, HTTPException, UploadFile, File, status

from api.models.ingestion import (
    IngestionErrorResponse,
    JDIngestionResponse,
    ResumeIngestionResponse,
)
from utils.logging import get_logger

logger = get_logger(__name__)

router = APIRouter(prefix="/ingest", tags=["ingestion"])

# ---------------------------------------------------------------------------
# File reading helpers
# ---------------------------------------------------------------------------

_ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_EXTENSION_MIME_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


async def _read_upload(file: UploadFile) -> tuple[bytes, str, str]:
    """
    Read and validate an uploaded file.

    Returns
    -------
    (content_bytes, filename, mime_type)
    """
    filename = file.filename or "unknown"
    content_type = file.content_type or ""

    # Infer MIME from extension if content_type is missing/generic
    if content_type not in _ALLOWED_CONTENT_TYPES:
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        content_type = _EXTENSION_MIME_MAP.get(ext, content_type)

    if content_type not in _ALLOWED_CONTENT_TYPES:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported file type: {content_type}. Only PDF and DOCX are accepted.",
        )

    content = await file.read()
    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty.",
        )

    return content, filename, content_type


async def _extract_text(content: bytes, mime_type: str) -> str:
    """
    Extract plain text from a PDF or DOCX file.
    """
    if mime_type == "application/pdf":
        try:
            from pypdf import PdfReader
            import io

            reader = PdfReader(io.BytesIO(content))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n\n".join(text_parts)
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to parse PDF: {exc}",
            )

    elif (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        try:
            from docx import Document
            import io

            doc = Document(io.BytesIO(content))
            text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
            return "\n\n".join(text_parts)
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to parse DOCX: {exc}",
            )

    raise HTTPException(
        status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
        detail=f"Cannot extract text from {mime_type}",
    )


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------


@router.post(
    "/resume",
    response_model=ResumeIngestionResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Ingest a resume",
    description="Upload a PDF or DOCX resume to be chunked, embedded, and indexed.",
)
async def ingest_resume(file: UploadFile = File(...)):
    """
    Ingest a single resume document.

    1. Parse PDF/DOCX to text
    2. Extract metadata (name, skills, experience, education)
    3. Chunk by sections
    4. Generate embeddings
    5. Store in the resumes Qdrant collection
    """
    content, filename, mime_type = await _read_upload(file)
    text = await _extract_text(content, mime_type)

    try:
        from core.config import DocumentType, get_settings
        from ingestion.interfaces import RawDocument
        from ingestion.processors.resume_metadata_extractor import ResumeMetadataExtractor
        from ingestion.processors.section_chunker import ResumeSectionChunker
        from ingestion.processors.cleaner import DefaultTextCleaner
        from ingestion.processors.validator import DefaultDocumentValidator
        from ingestion.embedders.factory import create_embedder
        from ingestion.pipeline import IngestionPipeline
        from utils.helpers import generate_document_id

        settings = get_settings()

        document = RawDocument(
            source=filename,
            content=text,
            mime_type=mime_type,
            document_type=DocumentType.RESUME,
        )

        # Build pipeline components
        validator = DefaultDocumentValidator(settings.ingestion)
        cleaner = DefaultTextCleaner()
        chunker = ResumeSectionChunker(settings.ingestion.chunking)
        extractor = ResumeMetadataExtractor()
        embedder = create_embedder(settings.ingestion.embedding)

        pipeline = IngestionPipeline(
            config=settings.ingestion,
            validator=validator,
            cleaner=cleaner,
            chunker=chunker,
            metadata_extractor=extractor,
            embedder=embedder,
        )

        result = await pipeline.ingest(document)

        if not result.success:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=result.error or "Unknown ingestion error",
            )

        return ResumeIngestionResponse(
            document_id=result.document_id,
            source=result.source,
            chunks_created=result.chunks_created,
            metadata=result.metadata,
        )

    except HTTPException:
        raise
    except Exception as exc:
        logger.error("Resume ingestion failed", error=str(exc))
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Internal error during resume ingestion: {exc}",
        )


@router.post(
    "/jd",
    response_model=JDIngestionResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Ingest a job description",
    description="Upload a PDF or DOCX job description to be chunked, embedded, and indexed.",
)
async def ingest_jd(file: UploadFile = File(...)):
    """
    Ingest a single job description document.

    1. Parse PDF/DOCX to text
    2. Extract metadata (job title, required skills, experience, education)
    3. Chunk by sections
    4. Generate embeddings
    5. Store in the job_descriptions Qdrant collection
    """
    content, filename, mime_type = await _read_upload(file)
    text = await _extract_text(content, mime_type)

    try:
        from core.config import DocumentType, get_settings
        from ingestion.interfaces import RawDocument
        from ingestion.processors.jd_metadata_extractor import JDMetadataExtractor
        from ingestion.processors.section_chunker import JDSectionChunker
        from ingestion.processors.cleaner import DefaultTextCleaner
        from ingestion.processors.validator import DefaultDocumentValidator
        from ingestion.embedders.factory import create_embedder
        from ingestion.pipeline import IngestionPipeline

        settings = get_settings()

        document = RawDocument(
            source=filename,
            content=text,
            mime_type=mime_type,
            document_type=DocumentType.JOB_DESCRIPTION,
        )

        validator = DefaultDocumentValidator(settings.ingestion)
        cleaner = DefaultTextCleaner()
        chunker = JDSectionChunker(settings.ingestion.chunking)
        extractor = JDMetadataExtractor()
        embedder = create_embedder(settings.ingestion.embedding)

        pipeline = IngestionPipeline(
            config=settings.ingestion,
            validator=validator,
            cleaner=cleaner,
            chunker=chunker,
            metadata_extractor=extractor,
            embedder=embedder,
        )

        result = await pipeline.ingest(document)

        if not result.success:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=result.error or "Unknown ingestion error",
            )

        return JDIngestionResponse(
            document_id=result.document_id,
            source=result.source,
            chunks_created=result.chunks_created,
            metadata=result.metadata,
        )

    except HTTPException:
        raise
    except Exception as exc:
        logger.error("JD ingestion failed", error=str(exc))
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Internal error during JD ingestion: {exc}",
        )
